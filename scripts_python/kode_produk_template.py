"""Buat template 'KODE PRODUK TANPA NAMA' agar user bisa mengisi nama produk untuk kode yang namanya kosong di Juald.
Isi kolom NAMA PRODUK lalu jalankan ulang build_dataset.py: nama itu otomatis dipakai."""
import datetime
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl
from openpyxl import Workbook

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
SRC = os.path.join(BASE, "2- DATA LAMA", "SUMBER AWAL", "BACK UP zaineva hijab_data - 29.08 SORE.xlsx")
OUT = os.path.join(BASE, "4- ISIAN TIM", "KODE PRODUK TANPA NAMA - isi nama di kolom terakhir.xlsx")
PCS = re.compile(r"\s*:\s*\d+\s*\(\s*pcs\s*\)\s*$", re.I)


def to_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, (int, float)):
        return (datetime.datetime(1899, 12, 30) + datetime.timedelta(days=v)).date()
    return None


wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
header = {r[1]: dict(status=r[23], plat=str(r[21] or ""), isi=str(r[10] or "")) for r in wb["Jual"].iter_rows(min_row=2, values_only=True)}

info = defaultdict(lambda: dict(n=0, rev=0.0, d0=None, d1=None, plat=Counter(), orders=[]))
for r in wb["Juald"].iter_rows(min_row=2, values_only=True):
    code = str(r[2]).strip() if r[2] is not None else ""
    name = PCS.sub("", str(r[6] or "")).strip()
    if not code or (name and not re.fullmatch(r"[\d.,\s-]*", name)):
        continue
    h = header.get(r[1])
    if not h or h["status"] != "Terkirim":
        continue
    x = info[code]
    x["n"] += 1
    x["rev"] += float(r[5] or 0)
    d = to_date(r[7])
    if d:
        x["d0"] = min(x["d0"], d) if x["d0"] else d
        x["d1"] = max(x["d1"], d) if x["d1"] else d
    x["plat"][h["plat"]] += 1
    if len(x["orders"]) < 3:
        x["orders"].append(str(r[1]))

out = Workbook()
ws = out.active
ws.title = "Kode Produk"
ws.append(["Kode Produk", "Jumlah Baris", "Omzet (Rp)", "Tanggal Pertama", "Tanggal Terakhir",
           "Platform Tersering", "Contoh No Pesanan", "NAMA PRODUK (ISI DI SINI)"])
for code, x in sorted(info.items(), key=lambda kv: -kv[1]["rev"]):
    ws.append([code, x["n"], x["rev"], x["d0"], x["d1"], x["plat"].most_common(1)[0][0],
               ", ".join(x["orders"]), None])
for col, w in zip("ABCDEFGH", [26, 13, 16, 15, 15, 20, 44, 42]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
note = out.create_sheet("Petunjuk")
for line in [
    "Sheet 'Kode Produk' berisi kode produk yang namanya kosong di data sumber (sheet Juald), urut omzet terbesar.",
    "Isi kolom H (NAMA PRODUK) dengan nama produk yang benar. Kode yang belum diisi dibiarkan kosong.",
    "Catatan: kode numerik panjang (19 digit) adalah ID produk marketplace; Excel hanya menyimpan 15 digit pertama dengan presisi,",
    "jadi 3-4 digit terakhir bisa berbeda dari ID aslinya. Cocokkan dengan 15 digit pertama.",
    "Setelah diisi, simpan file ini di folder yang sama lalu minta jalankan ulang ETL; nama akan otomatis terpakai.",
]:
    note.append([line])
note.column_dimensions["A"].width = 130
out.save(OUT)
print("kode:", len(info), "| baris:", sum(x["n"] for x in info.values()), "| omzet Rp", f"{sum(x['rev'] for x in info.values()):,.0f}")
print(OUT)
