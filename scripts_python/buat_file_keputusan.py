"""Satu file xlsx untuk tim: (1) nama produk ambigu, (2) pesanan yang kecamatannya kosong.
Kolom 'FINAL SKU' dan 'Kecamatan FINAL' otomatis dibaca ETL (file BERESIN*.xlsx di folder RAPIHKAN DATA).
Isian yang sudah ada di file lama (BERESIN-2 keputusan produk...) dibawa ke file baru, lalu file lama diarsipkan."""
import glob
import os
import shutil
import sys
from collections import defaultdict

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(__file__))
import build_dataset as b  # noqa: E402

sys.stdout.reconfigure(encoding="utf-8")
OUT = os.path.join(b.ISIAN_DIR, "BERESIN-2 produk ambigu + kecamatan kosong (isi kolom FINAL).xlsx")
AUDIT = b.AUDIT_XLSX
OLD = os.path.join(b.ISIAN_DIR, "BERESIN-2 keputusan produk (isi kolom FINAL SKU).xlsx")

b.load_master()

# --- isian lama (dibawa ke file baru)
prefill = {}
for path in (OLD, OUT):
    if os.path.exists(path):
        wb_old = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for ws in wb_old:
            name_i = fin_i = None
            for r in ws.iter_rows(values_only=True):
                if name_i is None:
                    cells = [b.clean_text(c).casefold() for c in r]
                    if "nama di data" in cells:
                        name_i = cells.index("nama di data")
                        fin_i = next((i for i, c in enumerate(cells) if "final" in c), None)
                    continue
                if fin_i is not None and len(r) > fin_i and r[name_i] and r[fin_i] and b.clean_text(r[fin_i]):
                    prefill[b.product_key(str(r[name_i]))] = b.clean_text(r[fin_i])

# --- sheet 1: produk ambigu
wb_in = openpyxl.load_workbook(b.ALIAS_FILE, read_only=True, data_only=True)
prod = []
for r in wb_in["Perlu keputusan"].iter_rows(min_row=2, values_only=True):
    if r[0] and not r[5]:
        prod.append((r[0], r[1], float(r[2] or 0), r[3] or "", r[4] or ""))
prod.sort(key=lambda x: -x[2])
prod_total = sum(x[2] for x in prod) or 1

wb = Workbook()
ws = wb.active
ws.title = "Produk Ambigu"
ws.append(["HANYA YANG BENAR-BENAR AMBIGU. Yang sudah yakin (singkatan, urutan kata, BD=Bergo Diamond, Jum=Jumbo) dan produk baru sudah diputuskan otomatis."])
ws.append(["Isi kolom FINAL SKU (dropdown SKU master; boleh ketik nama baru). Kosongkan kalau produknya memang produk baru. Urut omzet terbesar dulu."])
ws.append([])
ws.append(["No", "Nama di data", "Baris", "Omzet (Rp)", "% kumulatif", "Usulan sistem", "Alasan ambigu", "FINAL SKU (isi di sini)"])
H = 4
for c in ws[H]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
cum = 0.0
for i, (name, n, rev, sug, why) in enumerate(prod, 1):
    cum += rev
    ws.append([i, name, n, rev, round(cum / prod_total, 4), sug, why, prefill.get(b.product_key(str(name)))])
for col, w in zip("ABCDEFGH", [6, 40, 8, 15, 12, 62, 44, 42]):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f"C{H + 1}"
ws.auto_filter.ref = f"A{H}:H{H + len(prod)}"
for r in range(H + 1, H + len(prod) + 1):
    ws.cell(r, 4).number_format = "#,##0"
    ws.cell(r, 5).number_format = "0.0%"
    ws.cell(r, 8).fill = PatternFill("solid", fgColor="FFF2CC")

ms = wb.create_sheet("Master SKU")
ms.append(["SKU master", "Kategori", "Brand"])
for e in sorted(b.MASTER.values(), key=lambda e: e["sku"].casefold()):
    ms.append([e["sku"], e["cat"], e["brand"]])
ms.column_dimensions["A"].width = 46
ms.column_dimensions["B"].width = 24
dv = DataValidation(type="list", formula1=f"='Master SKU'!$A$2:$A${len(b.MASTER) + 1}", allow_blank=True, showErrorMessage=False)
ws.add_data_validation(dv)
dv.add(f"H{H + 1}:H{H + len(prod)}")

# --- sheet 2: kecamatan kosong (satu baris per pesanan)
orders = {}
wa = openpyxl.load_workbook(AUDIT, read_only=True, data_only=True)
for r in wa["Kecamatan Kosong"].iter_rows(min_row=2, values_only=True):
    no = r[1]
    if not no:
        continue
    o = orders.get(no)
    if o is None:
        orders[no] = dict(no=no, date=r[0], src=r[2], ch=r[4], city=r[13], prov=r[17], addr=r[18] or "",
                          why=r[20] or "", rev=float(r[12] or 0), lines=1)
    else:
        o["rev"] += float(r[12] or 0)
        o["lines"] += 1


def readable(o):
    return "disensor" not in o["why"].casefold() and "kosong di data" not in o["why"].casefold()


ords = sorted(orders.values(), key=lambda o: (not readable(o), -o["rev"]))
wk = wb.create_sheet("Kecamatan Kosong", 1)
wk.append(["PESANAN YANG KOTANYA ADA TAPI KECAMATANNYA KOSONG (satu baris per pesanan). Baris teratas: alamat terbaca, tim bisa mengisi kecamatan dari alamat."])
wk.append(["Isi kolom Kecamatan FINAL (huruf besar/kecil bebas). 'Alamat disensor' = marketplace menyembunyikan alamat, biasanya tidak bisa diisi. Kosongkan yang tidak diketahui."])
wk.append([])
wk.append(["No", "No Pesanan", "Tanggal", "Sumber", "Channel", "Kota", "Provinsi", "Baris", "Omzet (Rp)", "Alamat (dipotong)", "Status alamat", "Alasan", "Kecamatan FINAL (isi di sini)"])
for c in wk[H]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
for i, o in enumerate(ords, 1):
    wk.append([i, o["no"], o["date"], o["src"], o["ch"], o["city"], o["prov"], o["lines"], o["rev"], o["addr"],
               "Alamat terbaca" if readable(o) else "Alamat disensor / kosong", o["why"], None])
for col, w in zip("ABCDEFGHIJKLM", [7, 26, 12, 11, 12, 24, 20, 7, 14, 60, 22, 46, 34]):
    wk.column_dimensions[col].width = w
wk.freeze_panes = f"C{H + 1}"
wk.auto_filter.ref = f"A{H}:M{H + len(ords)}"
for r in range(H + 1, H + len(ords) + 1):
    wk.cell(r, 9).number_format = "#,##0"
    wk.cell(r, 13).fill = PatternFill("solid", fgColor="FFF2CC")

wb.save(OUT)

# arsipkan file lama (isiannya sudah dibawa)
archived = False
if os.path.exists(OLD):
    try:
        os.makedirs(os.path.join(b.BASE, "6- BISA DIBUANG"), exist_ok=True)
        shutil.move(OLD, os.path.join(b.BASE, "6- BISA DIBUANG", os.path.basename(OLD)))
        archived = True
    except PermissionError:
        print("CATATAN: file lama sedang terbuka di Excel, tidak dipindahkan ke 6- BISA DIBUANG (tutup dulu lalu jalankan ulang).")

n_read = sum(1 for o in ords if readable(o))
print("file:", OUT)
print("produk ambigu:", len(prod), "| omzet Rp", f"{sum(x[2] for x in prod):,.0f}", "| isian lama dibawa:", sum(1 for x in prod if b.product_key(str(x[0])) in prefill))
print("kecamatan kosong: pesanan", len(ords), "| alamat terbaca (bisa diisi):", n_read, "| omzet terbaca Rp", f"{sum(o['rev'] for o in ords if readable(o)):,.0f}",
      "| disensor/kosong:", len(ords) - n_read, "| omzet Rp", f"{sum(o['rev'] for o in ords if not readable(o)):,.0f}")
